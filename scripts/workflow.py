"""Build numbered reading bundles and import them using Zotero's local Connector."""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import html
import json
import re
import shutil
import sys
import unicodedata
import urllib.error
import urllib.request
import uuid
from pathlib import Path


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def write_json(path, value):
    path = Path(path)
    temp = path.with_name(path.name + ".writing-" + uuid.uuid4().hex)
    try:
        temp.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
        temp.replace(path)
    finally:
        temp.unlink(missing_ok=True)


def chinese_number(n):
    if not isinstance(n, int) or isinstance(n, bool) or not 1 <= n <= 9999:
        raise ValueError("number must be an integer from 1 to 9999")
    digits = "零一二三四五六七八九"
    out = ""
    gap = False
    for power, unit in [(1000, "千"), (100, "百"), (10, "十"), (1, "")]:
        digit, n = divmod(n, power)
        if digit:
            if gap:
                out += "零"
            out += digits[digit] + unit
            gap = False
        elif out and n:
            gap = True
    return out[1:] if out.startswith("一十") else out


def safe_name(value):
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).strip(" .")[:100].rstrip(" .")
    if not value:
        raise ValueError("Empty filename")
    if re.match(r"^(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\.|$)", value, re.I):
        value = "_" + value
    return value


def identity(paper):
    doi = str(paper.get("doi", "")).strip().lower()
    doi = re.sub(r"^https?://(?:dx\.)?doi\.org/", "", doi)
    if doi and not re.fullmatch(r"10\.\d{4,9}/\S+", doi):
        raise ValueError("Invalid DOI")
    title = re.sub(r"\s+", " ", unicodedata.normalize("NFKC", paper["title"])).strip().casefold()
    return doi, title


def child_dir(root, name):
    child = (root / name).resolve()
    if child == root or not child.is_relative_to(root):
        raise ValueError("Output subdirectory must stay inside paper_root")
    return child


def validate_pdf(path):
    from pypdf import PdfReader
    path = Path(path)
    with path.open("rb") as f:
        if f.read(5) != b"%PDF-":
            raise ValueError("Input is not a PDF; HTML/login pages are not accepted")
    reader = PdfReader(path)
    if reader.is_encrypted or not len(reader.pages):
        raise ValueError("PDF must be readable and have at least one page")
    return len(reader.pages), hashlib.sha256(path.read_bytes()).hexdigest()


def make_docx(path, paper, label):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml.ns import qn
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    sec.left_margin = sec.right_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.space_after = Pt(6)
    for name in ["Title", "Heading 1"]:
        style = doc.styles[name]
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "黑体")
    doc.add_paragraph(paper["title_zh"], style="Title")
    doc.add_paragraph(f"{paper['date']}  {label}  论文简读")
    doc.add_heading("论文信息", level=1)
    doc.add_paragraph(f"{paper['title']}\n{paper['journal']} ({paper['year']})\n"
                      f"DOI: {paper.get('doi', '')}\n正文来源: {paper['pdf_source']}")
    for section in paper["sections"]:
        doc.add_heading(section["heading"], level=1)
        doc.add_paragraph(section["text"])
    for root in [doc._element, doc.styles._element]:
        for border in root.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    doc.save(path)


def prepare(config_path, paper_path, pdf_path, work_path):
    config_path = Path(config_path).resolve()
    config, paper = read_json(config_path), read_json(paper_path)
    for key in ["title", "title_zh", "journal", "pdf_source", "url"]:
        if not isinstance(paper.get(key), str) or not paper[key].strip():
            raise ValueError(f"Missing {key}")
    dt.date.fromisoformat(paper["date"])
    if not isinstance(paper.get("year"), int) or not 1500 <= paper["year"] <= 3000:
        raise ValueError("Invalid publication year")
    if not paper.get("sections") or any(not x.get("heading") or not x.get("text") for x in paper["sections"]):
        raise ValueError("Reviewed reading sections are required")
    label = "文献" + chinese_number(paper["number"])
    doi, title = identity(paper)
    paper["doi"] = doi
    pages, digest = validate_pdf(pdf_path)
    root = (config_path.parent / Path(config["paper_root"])).resolve()
    pdf_dir = child_dir(root, config.get("pdf_subdir", "文献"))
    word_dir = child_dir(root, config.get("word_subdir", "简读"))
    collection = config.get("collection_name", "").strip()
    if not collection:
        raise ValueError("collection_name must name an existing Zotero collection")
    state_dir = root / ".paper-workflow"
    state_dir.mkdir(parents=True, exist_ok=True)
    lock = state_dir / "prepare.lock"
    with lock.open("x", encoding="utf-8") as handle:
        handle.write("Remove only after confirming no prepare process is running")
    created = []
    try:
        ledger_path = state_dir / "ledger.json"
        ledger = read_json(ledger_path) if ledger_path.exists() else []
        for row in ledger:
            if (doi and row.get("doi") == doi) or row["title_key"] == title:
                raise ValueError("Paper already exists in local ledger; inspect existing task")
            if row["number"] == paper["number"]:
                raise ValueError("Literature number already reserved")
        for directory in [pdf_dir, word_dir]:
            if directory.exists() and any(re.search(re.escape(label) + r"(?=[_｜.\s-]|$)", f.name)
                                          for f in directory.iterdir() if f.is_file()):
                raise ValueError("Existing file already uses this literature number")
        work = Path(work_path).resolve()
        work.mkdir(parents=True, exist_ok=True)
        if any(work.iterdir()):
            raise ValueError("Use an empty task directory; existing tasks must not be overwritten")
        pdf_dir.mkdir(parents=True, exist_ok=True)
        word_dir.mkdir(parents=True, exist_ok=True)
        prefix = f"{paper['date']}_{label}_"
        pdf_out = pdf_dir / (prefix + safe_name(paper["title"]) + ".pdf")
        word_out = word_dir / (prefix + safe_name(paper["title_zh"]) + ".docx")
        title_note = f"{label}｜要点｜{paper['title_zh']}（{paper['journal']} {paper['year']}）"
        note = "<h1>" + html.escape(title_note) + "</h1>"
        note += "<p>整理日期：" + html.escape(paper["date"]) + "</p>"
        note += "<p>DOI：" + html.escape(doi) + "；正文来源：" + html.escape(paper["pdf_source"]) + "</p>"
        for section in paper["sections"]:
            note += "<h2>" + html.escape(section["heading"]) + "</h2><p>" + html.escape(section["text"]).replace("\n", "<br/>") + "</p>"
        item = {
            "id": "paper-" + uuid.uuid4().hex, "itemType": "journalArticle",
            "title": paper["title"], "publicationTitle": paper["journal"],
            "date": str(paper["year"]), "DOI": doi, "url": paper["url"],
            "creators": [{"firstName": a.get("firstName", ""), "lastName": a["lastName"],
                          "creatorType": "author"} for a in paper.get("authors", [])],
            "tags": [{"tag": x} for x in dict.fromkeys([label, "每日精选", paper["date"], *paper.get("tags", [])])],
            "notes": [{"note": note}], "attachments": []
        }
        for key in ["volume", "issue", "pages"]:
            if key in paper:
                item[key] = str(paper[key])
        staged_doc = work / "preview.docx"
        make_docx(staged_doc, paper, label)
        for src, dst in [(Path(pdf_path), pdf_out), (staged_doc, word_out)]:
            with dst.open("xb") as target, src.open("rb") as source:
                created.append(dst)
                shutil.copyfileobj(source, target)
        write_json(work / "zotero.json", item)
        manifest = {"number": paper["number"], "label": label, "doi": doi,
                    "title_key": title, "note_title": title_note, "pdf_path": str(pdf_out),
                    "word_path": str(word_out), "pdf_pages": pages, "pdf_sha256": digest,
                    "collection_name": collection, "session_id": "paper-" + uuid.uuid4().hex,
                    "stage": "prepared", "visual_review": "pending", "zotero_verification": "pending"}
        write_json(work / "manifest.json", manifest)
        ledger.append({"number": paper["number"], "doi": doi, "title_key": title, "task": str(work)})
        write_json(ledger_path, ledger)
        return manifest
    except Exception:
        for path in created:
            path.unlink(missing_ok=True)
        raise
    finally:
        lock.unlink(missing_ok=True)


def connector(endpoint, payload=None, *, pdf=None, metadata=None):
    """Only local Zotero; never read credentials or inherit a proxy for loopback."""
    headers = {"Content-Type": "application/json"}
    body = json.dumps(payload or {}, ensure_ascii=False).encode("utf-8")
    if pdf is not None:
        headers = {"Content-Type": "application/pdf", "X-Metadata": json.dumps(metadata)}
        body = Path(pdf).read_bytes()
    request = urllib.request.Request("http://127.0.0.1:23119/connector/" + endpoint,
                                     data=body, headers=headers, method="POST")
    opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
    with opener.open(request, timeout=45) as response:
        raw = response.read()
        if response.status not in (200, 201, 204):
            raise RuntimeError(f"Unexpected Zotero status: {response.status}")
        return json.loads(raw) if raw else {}


def import_zotero(work_path, commit=False):
    work = Path(work_path)
    state_path = work / "manifest.json"
    state, item = read_json(state_path), read_json(work / "zotero.json")
    if state["stage"] == "complete":
        return {"stage": "complete", "message": "Already imported; verify in Zotero, no repeat write"}
    if state["stage"] in ["saving", "attaching", "uncertain"]:
        raise RuntimeError("Previous operation outcome is uncertain; inspect Zotero before retrying")
    pages, digest = validate_pdf(state["pdf_path"])
    if digest != state["pdf_sha256"]:
        raise ValueError("PDF changed since prepare; import cancelled")
    selected = connector("getSelectedCollection")
    matches = [target for target in selected.get("targets", [])
               if target.get("name") == state["collection_name"] and target.get("id", "").startswith("C")]
    if len(matches) != 1 or not matches[0].get("filesEditable"):
        raise ValueError("Collection must exist uniquely and allow files; inspect Zotero")
    target = matches[0]["id"]
    if not commit:
        return {"stage": "preview", "collection": state["collection_name"],
                "note_title": state["note_title"], "pdf_pages": pages, "writes": False}
    session = state["session_id"]
    if state["stage"] == "prepared":
        state["stage"] = "saving"
        write_json(state_path, state)
        try:
            connector("saveItems", {"sessionID": session, "items": [item], "uri": item["url"]})
        except Exception:
            state["stage"] = "uncertain"
            write_json(state_path, state)
            raise
        state["stage"] = "item_saved"
        write_json(state_path, state)
    if state["stage"] == "item_saved":
        connector("updateSession", {"sessionID": session, "target": target})
        state["stage"] = "classified"
        write_json(state_path, state)
    if state["stage"] == "classified":
        state["stage"] = "attaching"
        write_json(state_path, state)
        try:
            connector("saveAttachment", pdf=state["pdf_path"], metadata={
                "sessionID": session, "parentItemID": item["id"],
                "url": item["url"], "title": state["label"] + " 正文 PDF"})
        except Exception:
            state["stage"] = "uncertain"
            write_json(state_path, state)
            raise
        state["stage"] = "complete"
        write_json(state_path, state)
    return {"stage": state["stage"], "zotero_verification": "pending",
            "message": "Verify title, note, collection and attachment in Zotero"}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    build = commands.add_parser("prepare")
    for name in ["config", "paper", "pdf", "work"]:
        build.add_argument("--" + name, required=True)
    save = commands.add_parser("zotero-import")
    save.add_argument("--work", required=True)
    save.add_argument("--commit", action="store_true")
    args = parser.parse_args()
    try:
        if args.command == "prepare":
            result = prepare(args.config, args.paper, args.pdf, args.work)
        else:
            result = import_zotero(args.work, args.commit)
        print(json.dumps(result, ensure_ascii=False, indent=2))
    except Exception as exc:
        print(f"{type(exc).__name__}: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
